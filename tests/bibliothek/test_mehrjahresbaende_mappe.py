"""Die Exceldatei der Mehrjahresbände: schreiben, wiederlesen, eine Zelle ändern.

Der Punkt dieser Tests ist die **Rückfallebene**: die Datei muss ohne das
Dashboard lesbar bleiben und beim Öffnen wie die von Hand gepflegte Fassung
aussehen. Deshalb wird hier nicht nur der Inhalt geprüft, sondern auch der
Aufbau - Zellverbünde, Rahmen und die leere Ecke oben links.
"""
from __future__ import annotations

from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from mehrjahresbaende.core import (
    Jahrgangszeile,
    MappeUnlesbar,
    Sonderfall,
    Spalte,
    Uebersicht,
    Zelle,
    blatt,
    erlaubte_marken,
    lies_uebersicht,
    schreibe_datei,
    setze_marke,
)

SPALTEN = (
    Spalte("Deutsch", "Aufgabenfeld A"),
    Spalte("Englisch", "Aufgabenfeld A"),
    Spalte("Politik", "Aufgabenfeld B"),
    Spalte("Theater", "(ohne Aufgabenfeld)"),
)


def zeile(jahrgang: int, *marken: str) -> Jahrgangszeile:
    return Jahrgangszeile(
        jahrgang=jahrgang,
        zellen=tuple(Zelle(fach=spalte.fach, marke=marke)
                     for spalte, marke in zip(SPALTEN, marken)),
    )


@pytest.fixture()
def uebersicht() -> Uebersicht:
    return Uebersicht(
        spalten=SPALTEN,
        zeilen=(
            zeile(5, "X", "---", "---", ""),
            zeile(6, "", "X", "A", "B"),
            Jahrgangszeile(jahrgang=7, hinweis="Alle Bücher ohne erneute Anmeldung abgeben."),
        ),
        sonderfaelle=(Sonderfall("A", "nur „Politik 6“ muss abgegeben werden"),),
        schuljahr_alt="2025/2026", schuljahr_neu="2026/2027", erzeugt=date(2026, 9, 19),
    )


@pytest.fixture()
def datei(tmp_path: Path, uebersicht: Uebersicht) -> Path:
    pfad = tmp_path / "Mehrjahresbände Schulbuchausleihe.xlsx"
    schreibe_datei(pfad, uebersicht)
    return pfad


def test_geschrieben_und_wieder_gelesen_ist_dieselbe_uebersicht(
    datei: Path, uebersicht: Uebersicht,
) -> None:
    zurueck = lies_uebersicht(datei)
    assert zurueck.spalten == uebersicht.spalten
    assert zurueck.zeilen == uebersicht.zeilen
    assert zurueck.sonderfaelle == uebersicht.sonderfaelle
    assert (zurueck.schuljahr_alt, zurueck.schuljahr_neu) == ("2025/2026", "2026/2027")
    assert zurueck.erzeugt == date(2026, 9, 19)


def test_aufbau_entspricht_der_von_hand_gepflegten_fassung(datei: Path) -> None:
    ws = blatt(load_workbook(str(datei)))
    # Die Ecke oben links bleibt leer, hat aber ihren Rahmen.
    assert ws["A1"].value is None
    assert ws["A1"].border.left.style == "thick"
    # Aufgabenfelder als verbundene Kopfzellen über ihren Fächern.
    verbuende = {str(bereich) for bereich in ws.merged_cells.ranges}
    assert "B1:C1" in verbuende
    assert ws["B1"].value == "Aufgabenfeld A"
    assert ws["D1"].value == "Aufgabenfeld B"
    assert ws["E1"].value == "(ohne Aufgabenfeld)"
    # Fachnamen gedreht, Jahrgänge in Spalte A.
    assert ws["B2"].value == "Deutsch"
    assert ws["B2"].alignment.text_rotation == 90
    assert ws["A3"].value == "Jahrgang 5"
    # Die leere Marke ist eine leere Zelle, nicht der Text "".
    assert ws["E3"].value is None
    # Der Hinweis der individuellen Ausleihe läuft über alle Fachspalten.
    assert "B5:E5" in verbuende


def test_legende_steht_vollstaendig_in_der_datei(datei: Path) -> None:
    ws = blatt(load_workbook(str(datei)))
    texte = [ws.cell(nummer, 1).value for nummer in range(6, ws.max_row + 1)]
    assert "X = muss abgegeben werden" in texte
    assert any(str(text).startswith("A = nur „Politik 6“") for text in texte if text)
    assert any(str(text).startswith("Erzeugt am 19.09.2026") for text in texte if text)


def test_erlaubt_sind_die_festen_marken_plus_die_buchstaben_dieser_datei(
    datei: Path,
) -> None:
    assert erlaubte_marken(lies_uebersicht(datei)) == ("X", "", "---", "B", "A")


def test_setze_marke_findet_die_zelle_ueber_ihre_beschriftung(datei: Path) -> None:
    wb = load_workbook(str(datei))
    ws = blatt(wb)
    assert setze_marke(ws, 6, "Politik", "X") == "D4"
    assert ws["D4"].value == "X"
    # Leer heißt leere Zelle, nicht der leere Text.
    setze_marke(ws, 6, "Politik", "")
    assert ws["D4"].value is None


def test_setze_marke_lehnt_unbekannte_zeilen_und_spalten_ab(datei: Path) -> None:
    ws = blatt(load_workbook(str(datei)))
    with pytest.raises(MappeUnlesbar):
        setze_marke(ws, 99, "Deutsch", "X")
    with pytest.raises(MappeUnlesbar):
        setze_marke(ws, 5, "Chinesisch", "X")


def test_setze_marke_ruehrt_eine_hinweiszeile_nicht_an(datei: Path) -> None:
    """Jahrgang 7 leiht individuell aus - dort gibt es keine einzelnen Marken."""
    ws = blatt(load_workbook(str(datei)))
    with pytest.raises(MappeUnlesbar):
        setze_marke(ws, 7, "Deutsch", "X")


def test_neu_geschriebene_datei_ersetzt_die_alte_vollstaendig(
    datei: Path, uebersicht: Uebersicht,
) -> None:
    """Fällt eine Spalte weg, darf sie nicht als Rest stehen bleiben."""
    schmaler = Uebersicht(
        spalten=(SPALTEN[0],),
        zeilen=(Jahrgangszeile(jahrgang=5, zellen=(Zelle("Deutsch", "X"),)),),
        schuljahr_alt="2025/2026", schuljahr_neu="2026/2027", erzeugt=date(2026, 9, 19),
    )
    schreibe_datei(datei, schmaler)
    zurueck = lies_uebersicht(datei)
    assert [spalte.fach for spalte in zurueck.spalten] == ["Deutsch"]
    assert len(zurueck.zeilen) == 1
