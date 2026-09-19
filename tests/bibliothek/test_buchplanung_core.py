"""buchplanung/core: Mappe, Abgleich und die gerechneten Status - ohne Netz.

Die beiden Fragen, an denen diese Datei hängt:

1. **Überlebt Eingetragenes den Abgleich?** Eine Preisprüfung, eine Rücklage
   und eine Planungszeile müssen einen Abruf aus IServ überstehen - sonst wäre
   die Datei nach dem ersten „Aktualisieren" leer.
2. **Fällt auf, wenn sich etwas ändert?** Ein geänderter Preis muss die
   Prüfung kippen, ein neues Buch die Fachbestätigung veralten lassen.
"""
from __future__ import annotations

from datetime import date

import pytest

from buchplanung.core import (
    FACH_BESTAETIGT,
    FACH_OFFEN,
    FACH_VERALTET,
    NUR_PLANUNG,
    PLANUNG_AUSGEMUSTERT,
    PLANUNG_GEPLANT,
    PLANUNG_LAEUFT_AUS,
    PREIS_ABWEICHEND,
    PREIS_BESTAETIGT,
    PREIS_OFFEN,
    Buch,
    Buchplanung,
    Schnappschuss,
    UnbekanntesBuch,
    UngueltigeEingabe,
    fach_status,
    lade_schnappschuss,
    lies_datei,
    planungs_status,
    preis_status,
    schreibe_datei,
    setze_fachbestaetigung,
    setze_planung,
    setze_preis,
    setze_preise_des_verlags,
    setze_ruecklage,
    zusammenfuehren,
)
from buchplanung.core.modelle import Planungszeile

DEUTSCH = "9783060000005"
TERRA = "9783121000562"
ALT = "9783120000009"


def _buecher() -> tuple[Buch, ...]:
    return (
        Buch(isbn=DEUTSCH, titel="Deutschbuch 5", verlag="Cornelsen", faecher=("Deutsch",),
             jahrgaenge_vorjahr=(5,), jahrgaenge_aktuell=(5,), leihbar=True,
             neupreis=22.5, leihgebuehr=5.0),
        Buch(isbn=TERRA, titel="Terra 5/6", verlag="Klett", faecher=("Erdkunde", "Politik"),
             jahrgaenge_vorjahr=(5, 6), jahrgaenge_aktuell=(5, 6), leihbar=True,
             neupreis=25.0, leihgebuehr=5.0),
        # Nur im Vorjahr: das ist der Fall "ausgemustert", für den eine
        # Fachschaft Exemplare zurücklegen lassen möchte.
        Buch(isbn=ALT, titel="Chemie heute 9", verlag="Westermann", faecher=("Chemie",),
             jahrgaenge_vorjahr=(9,), leihbar=True, neupreis=30.0),
    )


def _schnappschuss(buecher=None, **felder) -> Schnappschuss:
    return Schnappschuss(
        schuljahr=felder.pop("schuljahr", "2026/2027"),
        vorjahr=felder.pop("vorjahr", "2025/2026"),
        buecher=buecher if buecher is not None else _buecher(),
        stand=felder.pop("stand", date(2026, 9, 19)),
        **felder,
    )


@pytest.fixture()
def stand() -> Buchplanung:
    return zusammenfuehren(None, _schnappschuss())


# ── Der Aufbau der Datei ─────────────────────────────────────────────────────


def test_die_drei_achsen_gruppieren_wie_erwartet(stand):
    # Jedes Buch hat genau einen Verlag; geprüft werden nur die des laufenden
    # Schuljahres - "Chemie heute 9" gibt es nur noch im Vorjahr.
    assert stand.verlage == ("Cornelsen", "Klett")
    # Ein Buch mit zwei Fächern steht in beiden.
    assert stand.faecher == ("Chemie", "Deutsch", "Erdkunde", "Politik")
    assert [b.isbn for b in stand.buecher_je_fach("Politik")] == [TERRA]
    assert stand.jahrgaenge == (5, 6, 9)


def test_schreiben_und_lesen_ergibt_denselben_stand(tmp_path, stand):
    stand = setze_preis(stand, isbn=DEUTSCH, preis=22.5, kuerzel="MLR",
                        datum=date(2026, 9, 1))
    stand = setze_ruecklage(stand, isbn=ALT, fach="Chemie", anzahl=5,
                            kuerzel="FK", datum=date(2026, 9, 2),
                            bemerkung="für die Sammlung")
    stand = setze_planung(stand, isbn=ALT, jahrgang=9, ausgemustert_nach="2026/2027",
                          beschluss="FK 12.05.2026")
    stand = setze_fachbestaetigung(stand, fach="Deutsch", kuerzel="ABC",
                                   datum=date(2026, 9, 3))

    pfad = tmp_path / "Buchplanung.xlsx"
    schreibe_datei(pfad, stand)
    gelesen = lies_datei(pfad)

    assert gelesen.schuljahr == "2026/2027"
    assert gelesen.vorjahr == "2025/2026"
    assert gelesen.stand == date(2026, 9, 19)
    assert {b.isbn for b in gelesen.buecher} == {DEUTSCH, TERRA, ALT}

    pruefung = gelesen.pruefung(DEUTSCH)
    assert (pruefung.preis, pruefung.kuerzel, pruefung.datum) == (22.5, "MLR", date(2026, 9, 1))
    ruecklage = gelesen.ruecklage(ALT, "Chemie")
    assert (ruecklage.anzahl, ruecklage.bemerkung) == (5, "für die Sammlung")
    zeile = gelesen.planungszeile(ALT, 9)
    assert (zeile.ausgemustert_nach, zeile.beschluss) == ("2026/2027", "FK 12.05.2026")
    assert gelesen.bestaetigung("Deutsch").bestaetigte_isbns == (DEUTSCH,)


def test_buecher_kommen_mit_jahrgaengen_und_herkunft_zurueck(tmp_path, stand):
    pfad = tmp_path / "Buchplanung.xlsx"
    schreibe_datei(pfad, stand)
    gelesen = lies_datei(pfad)

    terra = gelesen.buch(TERRA)
    assert terra.jahrgaenge_vorjahr == (5, 6)
    assert terra.jahrgaenge_aktuell == (5, 6)
    assert terra.faecher == ("Erdkunde", "Politik")
    assert terra.neupreis == 25.0

    alt = gelesen.buch(ALT)
    assert alt.jahrgaenge_aktuell == ()
    assert alt.jahrgaenge_vorjahr == (9,)
    # Nur im Vorjahr: steht nicht auf dem Verlagsblatt, hat hier also keinen Preis.
    assert alt.neupreis is None


def test_eine_planung_ohne_heutiges_vorkommen_bleibt_erhalten(tmp_path, stand):
    """Einführung in einen Jahrgang, der das Buch noch gar nicht führt."""
    stand = setze_planung(stand, isbn=DEUTSCH, jahrgang=7, eingefuehrt_ab="2028/2029")
    pfad = tmp_path / "Buchplanung.xlsx"
    schreibe_datei(pfad, stand)
    gelesen = lies_datei(pfad)

    zeile = gelesen.planungszeile(DEUTSCH, 7)
    assert zeile.eingefuehrt_ab == "2028/2029"
    # Die Zeile darf keinen Jahrgang erfinden, den IServ nie gemeldet hat.
    assert gelesen.buch(DEUTSCH).jahrgaenge == (5,)
    assert gelesen.herkunft(gelesen.buch(DEUTSCH), 7) == NUR_PLANUNG


def test_von_hand_geaenderter_mitgefuehrter_wert_wirkt_nicht_zurueck(tmp_path, stand):
    """Wer den Titel im Fach-Blatt überschreibt, ändert nichts an der Wahrheit."""
    from openpyxl import load_workbook

    pfad = tmp_path / "Buchplanung.xlsx"
    schreibe_datei(pfad, stand)

    wb = load_workbook(str(pfad))
    ws = wb["Bücher je Fach"]
    ws.cell(2, 2).value = "Von Hand verbogen"
    wb.save(str(pfad))

    # Der nächste Abgleich holt den Titel aus IServ zurück.
    neu = zusammenfuehren(lies_datei(pfad), _schnappschuss())
    assert "Von Hand verbogen" not in {b.titel for b in neu.buecher}


# ── Der Abgleich ─────────────────────────────────────────────────────────────


def test_eintragungen_ueberleben_den_abgleich(stand):
    stand = setze_preis(stand, isbn=DEUTSCH, preis=22.5, kuerzel="MLR", datum=date(2026, 9, 1))
    stand = setze_ruecklage(stand, isbn=ALT, fach="Chemie", anzahl=5)
    stand = setze_planung(stand, isbn=TERRA, jahrgang=7, eingefuehrt_ab="2027/2028")
    stand = setze_fachbestaetigung(stand, fach="Deutsch", kuerzel="ABC", datum=None)

    neu = zusammenfuehren(stand, _schnappschuss())

    assert neu.pruefung(DEUTSCH).preis == 22.5
    assert neu.ruecklage(ALT, "Chemie").anzahl == 5
    assert neu.planungszeile(TERRA, 7).eingefuehrt_ab == "2027/2028"
    assert neu.bestaetigung("Deutsch").kuerzel == "ABC"


def test_verschwundenes_buch_wird_verworfen_aber_gemeldet(stand):
    stand = setze_preis(stand, isbn=ALT, preis=30.0, kuerzel="MLR", datum=None)
    ohne_alt = tuple(b for b in _buecher() if b.isbn != ALT)

    neu = zusammenfuehren(stand, _schnappschuss(ohne_alt))

    assert neu.pruefung(ALT) is None
    assert any(ALT in warnung for warnung in neu.warnungen)


def test_fehlendes_vorjahr_ist_kein_fehler_sondern_eine_warnung(stand):
    neu = zusammenfuehren(stand, _schnappschuss(warnungen=("Das Vorjahr fehlt.",)))
    assert "Das Vorjahr fehlt." in neu.warnungen


# ── Die gerechneten Status ───────────────────────────────────────────────────


def test_preisstatus_kippt_wenn_iserv_den_preis_aendert(stand):
    stand = setze_preis(stand, isbn=DEUTSCH, preis=22.5, kuerzel="MLR", datum=None)
    buch = stand.buch(DEUTSCH)
    assert preis_status(buch, stand.pruefung(DEUTSCH))[0] == PREIS_BESTAETIGT

    teurer = tuple(
        b if b.isbn != DEUTSCH else Buch(**{**b.__dict__, "neupreis": 24.9})
        for b in _buecher()
    )
    neu = zusammenfuehren(stand, _schnappschuss(teurer))
    status, hinweis = preis_status(neu.buch(DEUTSCH), neu.pruefung(DEUTSCH))
    assert status == PREIS_ABWEICHEND
    assert "22,50" in hinweis and "24,90" in hinweis


def test_neues_buch_steht_von_allein_auf_offen(stand):
    stand = setze_preise_des_verlags(stand, verlag="Cornelsen", kuerzel="MLR", datum=None)[0]
    zusaetzlich = _buecher() + (
        Buch(isbn="9783060000012", titel="Deutschbuch 6", verlag="Cornelsen",
             faecher=("Deutsch",), jahrgaenge_aktuell=(6,), leihbar=True, neupreis=23.0),
    )
    neu = zusammenfuehren(stand, _schnappschuss(zusaetzlich))

    assert preis_status(neu.buch(DEUTSCH), neu.pruefung(DEUTSCH))[0] == PREIS_BESTAETIGT
    assert preis_status(neu.buch("9783060000012"), neu.pruefung("9783060000012"))[0] == PREIS_OFFEN
    assert any("geprüften Preis" in warnung for warnung in neu.warnungen)


def test_ganze_verlagsliste_auf_einmal(stand):
    neu, anzahl = setze_preise_des_verlags(stand, verlag="Klett", kuerzel="MLR",
                                           datum=date(2026, 9, 5))
    assert anzahl == 1
    assert neu.pruefung(TERRA).preis == 25.0
    assert neu.pruefung(TERRA).kuerzel == "MLR"
    assert neu.pruefung(DEUTSCH) is None


def test_fachbestaetigung_veraltet_mit_angabe_was_sich_geaendert_hat(stand):
    stand = setze_fachbestaetigung(stand, fach="Deutsch", kuerzel="ABC", datum=None)
    assert fach_status("Deutsch", stand.buecher_je_fach("Deutsch"),
                       stand.bestaetigung("Deutsch"))[0] == FACH_BESTAETIGT

    zusaetzlich = _buecher() + (
        Buch(isbn="9783060000012", titel="Deutschbuch 6", verlag="Cornelsen",
             faecher=("Deutsch",), jahrgaenge_aktuell=(6,), leihbar=True, neupreis=23.0),
    )
    neu = zusammenfuehren(stand, _schnappschuss(zusaetzlich))
    status, hinweis = fach_status("Deutsch", neu.buecher_je_fach("Deutsch"),
                                  neu.bestaetigung("Deutsch"))
    assert status == FACH_VERALTET
    assert "hinzugekommen: Deutschbuch 6" in hinweis


def test_fach_ohne_bestaetigung_ist_offen(stand):
    assert fach_status("Erdkunde", stand.buecher_je_fach("Erdkunde"), None)[0] == FACH_OFFEN


@pytest.mark.parametrize(("zeile", "erwartet"), [
    (Planungszeile(isbn=DEUTSCH, jahrgang=5, eingefuehrt_ab="2028/2029"), PLANUNG_GEPLANT),
    (Planungszeile(isbn=DEUTSCH, jahrgang=5, ausgemustert_nach="2026/2027"), PLANUNG_LAEUFT_AUS),
    (Planungszeile(isbn=DEUTSCH, jahrgang=5, ausgemustert_nach="2025/2026"), PLANUNG_AUSGEMUSTERT),
])
def test_planungsstatus_wird_gegen_das_schuljahr_gerechnet(zeile, erwartet):
    assert planungs_status(zeile, "2026/2027") == erwartet


# ── Was nicht eingetragen werden darf ────────────────────────────────────────


def test_unbekannte_isbn_wird_abgelehnt(stand):
    with pytest.raises(UnbekanntesBuch):
        setze_preis(stand, isbn="9780000000000", preis=1.0, kuerzel="MLR", datum=None)


def test_ausmusterung_vor_der_einfuehrung_wird_abgelehnt(stand):
    with pytest.raises(UngueltigeEingabe):
        setze_planung(stand, isbn=DEUTSCH, jahrgang=5,
                      eingefuehrt_ab="2027/2028", ausgemustert_nach="2026/2027")


def test_schuljahr_ohne_schraegstrich_wird_abgelehnt(stand):
    with pytest.raises(UngueltigeEingabe):
        setze_planung(stand, isbn=DEUTSCH, jahrgang=5, eingefuehrt_ab="2027")


def test_ruecklage_fuer_ein_fremdes_fach_wird_abgelehnt(stand):
    with pytest.raises(UngueltigeEingabe):
        setze_ruecklage(stand, isbn=DEUTSCH, fach="Chemie", anzahl=3)


def test_ruecklage_bekommt_ohne_angabe_den_status_gewuenscht(stand):
    neu = setze_ruecklage(stand, isbn=ALT, fach="Chemie", anzahl=4)
    assert neu.ruecklage(ALT, "Chemie").status == "gewünscht"


# ── Laden aus IServ ──────────────────────────────────────────────────────────


class _ZweiJahre:
    """Ein IServ, in dem sich das Vorjahr vom laufenden Jahr unterscheidet."""

    LISTEN = {
        "2026/2027": {5: [(DEUTSCH, "Deutschbuch 5", ["Deutsch"], 22.5)]},
        "2025/2026": {5: [(DEUTSCH, "Deutschbuch 5", ["Deutsch"], 21.0)],
                      9: [(ALT, "Chemie heute 9", ["Chemie"], 30.0)]},
    }

    class _Schuljahre:
        # Kennung und Anzeigename sind verschieden - wie im echten IServ.
        def get_current(self):
            return {"id": "2026/2027", "name": "Schuljahr 26/27"}

        def get_by_id(self, kennung):
            return {"id": kennung, "name": "Schuljahr 26/27"}

        def get_booklists(self, kennung):
            if kennung not in _ZweiJahre.LISTEN:
                raise LookupError(f"Schuljahr {kennung} gibt es nicht.")
            return [{"id": 100 + g, "grade": g} for g in sorted(_ZweiJahre.LISTEN[kennung])]

        def get_booklist(self, kennung, listen_id):
            grade = listen_id - 100
            return {"sections": [{"options": [{"items": [
                {"borrowable": True, "series": isbn,
                 "series_data": {"isbn": isbn, "title": titel, "subjectsFlat": faecher,
                                 "publisher": "Cornelsen", "price": preis, "fee": 5.0}}
                for isbn, titel, faecher, preis in _ZweiJahre.LISTEN[kennung][grade]
            ]}]}]}

    def __init__(self):
        self.schoolyears = self._Schuljahre()


def test_schnappschuss_legt_beide_schuljahre_uebereinander():
    schnappschuss = lade_schnappschuss(_ZweiJahre(), heute=date(2026, 9, 19))

    # Die Kennung, nicht der Name: mit ihr adressiert IServ das Schuljahr, und
    # nur sie lässt sich mit "eingeführt ab" und "ausgemustert nach" vergleichen.
    assert schnappschuss.schuljahr == "2026/2027"
    assert schnappschuss.name == "Schuljahr 26/27"
    assert schnappschuss.vorjahr == "2025/2026"
    je_isbn = {buch.isbn: buch for buch in schnappschuss.buecher}
    assert je_isbn[DEUTSCH].jahrgaenge_aktuell == (5,)
    assert je_isbn[DEUTSCH].jahrgaenge_vorjahr == (5,)
    # Nur im Vorjahr: ausgemustert.
    assert je_isbn[ALT].jahrgaenge_aktuell == ()
    assert je_isbn[ALT].jahrgaenge_vorjahr == (9,)
    assert not schnappschuss.warnungen


def test_fehlendes_vorjahr_bricht_den_abruf_nicht_ab():
    schnappschuss = lade_schnappschuss(_ZweiJahre(), vorjahr="2019/2020",
                                       heute=date(2026, 9, 19))
    assert {b.isbn for b in schnappschuss.buecher} == {DEUTSCH}
    assert any("2019/2020" in warnung for warnung in schnappschuss.warnungen)
